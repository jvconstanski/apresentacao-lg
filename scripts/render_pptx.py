"""deck.json → PPTX nativo e editável, na identidade visual da LG.

Uso:
    python render_pptx.py caminho/deck.json [-o saida.pptx]

Texto em caixas de texto reais, gráficos como charts nativos do PowerPoint (com as
cores da marca), tabela nativa, logos em PNG, rotas e paradas como formas
editáveis. Slide 16:9 (13,333 × 7,5 pol). Coordenadas desenhadas em px numa tela de
1920 × 1080 (1 px = 6.350 EMU = 0,5 pt), a mesma do HTML.

As fontes da marca (Archivo SemiCondensed, Archivo, IBM Plex Sans, IBM Plex Mono)
precisam estar instaladas em quem abre o arquivo. Os TTF estão em assets/fonts/ttf.
Sem elas, o PowerPoint troca por uma sem serifa do sistema (pitchFamily declarado).
"""
from __future__ import annotations

import argparse
import copy
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lg_comum as C  # noqa: E402
import render_html as RH  # noqa: E402

from lxml import etree  # noqa: E402
from PIL import ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

EMU_PX = 6350
MX, MT = 112, 72
BASE = 968          # limite inferior do conteúdo (rodapé começa em 1000)
FONTE_Y = 940       # linha "Fonte:" acima do rodapé

FONTES = {
    # chave: (nome no PowerPoint, negrito, arquivo TTF, espaçamento em em, pitchFamily)
    "display": ("Archivo SemiCondensed ExtraBold", False, "ArchivoSemiCondensed-ExtraBold.ttf", -0.03, "34"),
    "h3": ("Archivo SemiCondensed", True, "ArchivoSemiCondensed-Bold.ttf", -0.015, "34"),
    "rotulo": ("Archivo SemiBold", False, "Archivo-SemiBold.ttf", 0.1, "34"),
    "texto": ("IBM Plex Sans", False, "IBMPlexSans-Regular.ttf", 0.0, "34"),
    "texto-sb": ("IBM Plex Sans SmBld", False, "IBMPlexSans-SemiBold.ttf", 0.0, "34"),
    "mono": ("IBM Plex Mono", False, "IBMPlexMono-Regular.ttf", 0.0, "49"),
}
NOME_TEXTO = FONTES["texto"][0]


def E(px: float) -> Emu:
    return Emu(int(round(px * EMU_PX)))


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#").upper())


def pt(px: float) -> Pt:
    return Pt(round(px * 0.5 * 2) / 2)


# --------------------------------------------------------------------------
# medição de texto com as próprias fontes da marca
# --------------------------------------------------------------------------
class Medidor:
    def __init__(self):
        self._cache: dict = {}

    def fonte(self, estilo: str, px: float):
        chave = (estilo, int(round(px)))
        if chave not in self._cache:
            arq = C.FONTES_TTF / FONTES[estilo][2]
            self._cache[chave] = ImageFont.truetype(str(arq), max(1, int(round(px))))
        return self._cache[chave]

    def largura(self, texto: str, estilo: str, px: float, caps: bool = False) -> float:
        if caps:
            texto = texto.upper()
        f = self.fonte(estilo, px)
        w = f.getlength(texto) * (px / max(1, int(round(px))))
        return w + FONTES[estilo][3] * px * max(0, len(texto) - 1)

    def linhas(self, texto: str, estilo: str, px: float, largura: float, caps: bool = False) -> list[str]:
        saida = []
        for par in str(texto).split("\n"):
            palavras = [p for p in par.split(" ") if p]  # NBSP não quebra
            atual = ""
            for p in palavras:
                teste = f"{atual} {p}".strip()
                if not atual or self.largura(teste, estilo, px, caps) <= largura * 0.995:
                    atual = teste
                else:
                    saida.append(atual)
                    atual = p
            saida.append(atual)
        return saida

    def altura(self, texto: str, estilo: str, px: float, largura: float, lh: float, caps=False) -> float:
        return len(self.linhas(texto, estilo, px, largura, caps)) * px * lh


MED = Medidor()
NBSP = " "


def puro_nbsp(conteudo) -> str:
    """Texto puro em que a marca [DADO: ...] não quebra de linha."""
    return "".join(t.replace(" ", NBSP) if tipo == "dado" else t for t, tipo in C.segmentos(conteudo))


def altura_lista(w, itens, px=28, recuo=36, gap=16, lh=1.4):
    return sum(MED.altura(puro_nbsp(i), "texto", px, w - recuo, lh) for i in itens) + gap * max(0, len(itens) - 1)


# --------------------------------------------------------------------------
# construtor de slide
# --------------------------------------------------------------------------
class Folha:
    def __init__(self, prs, s: dict, n: int, ctx: dict):
        self.prs, self.s, self.n, self.ctx = prs, s, n, ctx
        self.sl = prs.slides.add_slide(prs.slide_layouts[6])
        self.sh = self.sl.shapes
        self.fundo = s.get("fundo") or C.LAYOUTS[s["layout"]]["fundo"]
        cor_fundo = {"claro": C.COR["fundo"], "apoio": C.COR["apoio"], "escuro": C.COR["grafite"], "azul": C.COR["azul"]}[self.fundo]
        self.cor_fundo = cor_fundo
        bg = self.sl.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(cor_fundo)
        escuro = self.fundo in ("escuro", "azul")
        self.c_txt = C.COR["fundo"] if self.fundo == "escuro" else ("#FFFFFF" if self.fundo == "azul" else C.COR["grafite"])
        self.c_apoio = C.COR["cinza_claro"] if self.fundo == "escuro" else ("#EAEDF0" if self.fundo == "azul" else C.COR["cinza_texto"])
        self.c_hl = C.COR["azul_claro"] if self.fundo == "escuro" else ("#FFFFFF" if self.fundo == "azul" else C.COR["azul"])
        self.escuro = escuro
        self.avisos: list[str] = []

    # ---- formas básicas -------------------------------------------------
    def ret(self, x, y, w, h, cor=None, linha=None, larg=0, forma=MSO_SHAPE.RECTANGLE):
        shp = self.sh.add_shape(forma, E(x), E(y), E(w), E(h))
        shp.shadow.inherit = False
        if cor:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(cor)
        else:
            shp.fill.background()
        if linha:
            shp.line.color.rgb = rgb(linha)
            shp.line.width = E(larg)
            etree.SubElement(shp.line._get_or_add_ln(), qn("a:miter"), lim="800000")
        else:
            shp.line.fill.background()
        return shp

    def poli(self, pontos, cor, larg, alpha=None, fechar=False, preencher=None):
        x0, y0 = pontos[0]
        ff = self.sh.build_freeform(x0, y0, scale=EMU_PX)
        ff.add_line_segments(pontos[1:], close=fechar)
        shp = ff.convert_to_shape(0, 0)
        shp.shadow.inherit = False
        if preencher:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(preencher)
        else:
            shp.fill.background()
        if cor:
            shp.line.color.rgb = rgb(cor)
            shp.line.width = E(larg)
            ln = shp.line._get_or_add_ln()
            if alpha is not None:
                clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
            etree.SubElement(ln, qn("a:miter"), lim="800000")
        else:
            shp.line.fill.background()
        return shp

    def quadrado(self, cx, cy, lado, cor, linha=None, larg=0, alpha=None):
        shp = self.ret(cx - lado / 2, cy - lado / 2, lado, lado, cor, linha, larg)
        if alpha is not None and linha:
            clr = shp.line._get_or_add_ln().find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
        return shp

    def tracado(self, x, y, h=40, w=340):
        k = h / 40
        pts = [(x, y + h), (x + 40 * k, y), (x + w, y), (x + w, y + 8.6 * k), (x + 43.6 * k, y + 8.6 * k), (x + 12.1 * k, y + h)]
        shp = self.poli(pts + [pts[0]], None, 0, fechar=True)
        shp.fill.gradient()
        shp.fill.gradient_angle = 0
        st = shp.fill.gradient_stops
        st[0].color.rgb = rgb(C.COR["azul"])
        st[0].position = 0.0
        st[1].color.rgb = rgb(C.COR["verde"])
        st[1].position = 1.0
        shp.name = "Traçado de conversão"
        return h

    def imagem(self, arq: Path, x, y, w, h=None):
        if h is None:
            return self.sh.add_picture(str(arq), E(x), E(y), width=E(w))
        return self.sh.add_picture(str(arq), E(x), E(y), width=E(w), height=E(h))

    def planta(self, variante: str, x, y, w, h, crop=(0, 0, 0, 0)):
        arq = C.GRAFISMO / {"claro": "planta.png", "escuro": "planta-escura.png", "apoio": "planta-apoio.png"}[variante]
        pic = self.imagem(arq, x, y, w, h)
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        pic.name = "Malha de quadras"
        return pic

    def esmaecer(self, x, y, w, h, cor):
        """Faixa em gradiente da cor de fundo para transparente (borda da planta)."""
        shp = self.ret(x, y, w, h, cor)
        shp.fill.gradient()
        shp.fill.gradient_angle = 0
        st = shp.fill.gradient_stops
        st[0].color.rgb = rgb(cor)
        st[0].position = 0.0
        st[1].color.rgb = rgb(cor)
        st[1].position = 1.0
        gs = shp.fill._xPr.find(qn("a:gradFill")).find(qn("a:gsLst")).findall(qn("a:gs"))
        etree.SubElement(gs[-1].find(qn("a:srgbClr")), qn("a:alpha"), val="0")
        return shp

    # ---- texto ----------------------------------------------------------
    def _run(self, p, texto, estilo, px, cor, tipo="normal", caps=False):
        r = p.add_run()
        r.text = texto
        chave = {"dado": "texto-sb", "codigo": "mono"}.get(tipo, estilo)
        nome, negrito, _, spc, pitch = FONTES[chave]
        f = r.font
        f.name = nome
        f.size = pt(px)
        f.bold = negrito
        cor_final = cor
        if tipo == "destaque":
            cor_final = self.c_hl
            if self.fundo == "azul":
                f.underline = True
        elif tipo == "dado":
            cor_final = "#EAEDF0" if self.escuro else C.COR["cinza_texto"]
        f.color.rgb = rgb(cor_final)
        rpr = r._r.get_or_add_rPr()
        if spc:
            rpr.set("spc", str(int(round(spc * px * 0.5 * 100))))
        if caps:
            rpr.set("cap", "all")
        latin = rpr.find(qn("a:latin"))
        if latin is not None:
            latin.set("pitchFamily", pitch)
            latin.set("charset", "0")
        if tipo == "dado":
            hl = etree.Element(qn("a:highlight"))
            etree.SubElement(hl, qn("a:srgbClr"), val=("3A414A" if self.escuro else "DCE3EC"))
            if latin is not None:
                latin.addprevious(hl)
            else:
                rpr.append(hl)
        return r

    def texto(self, x, y, w, conteudo, estilo="texto", px=28, cor=None, lh=1.4, caps=False,
              alinh="l", h=None, ancora="t", min_frac=0.8, max_linhas=None, nome=None, equilibrar=None):
        """Escreve um parágrafo (ou vários, separados por \\n). Retorna (altura, px usado)."""
        cor = cor or self.c_txt
        if equilibrar is None:
            equilibrar = estilo in ("display", "h3")
        puro = puro_nbsp(conteudo)
        px_f = px
        if h is not None or max_linhas:
            while px_f > px * min_frac:
                lin = MED.linhas(puro, estilo, px_f, w, caps)
                alt = len(lin) * px_f * lh
                cabe = (h is None or alt <= h + 0.5) and (not max_linhas or len(lin) <= max_linhas)
                if cabe:
                    break
                px_f *= 0.96
        lin = MED.linhas(puro, estilo, px_f, w, caps)
        alt = len(lin) * px_f * lh
        larg_caixa = w
        if equilibrar and len(lin) > 1 and "\n" not in str(conteudo):
            # mesmo número de linhas na menor largura possível (como text-wrap: balance)
            lo, hi = w * 0.5, w
            for _ in range(14):
                mid = (lo + hi) / 2
                if len(MED.linhas(puro, estilo, px_f, mid, caps)) <= len(lin):
                    hi = mid
                else:
                    lo = mid
            larg_caixa = min(w, hi * 1.02 + 4)
        if h is not None and alt > h + 1:
            self.avisos.append(f"slide {self.n:02d}: texto maior que a caixa ({puro[:50]}…)")
        caixa_h = h if h is not None else alt
        if alinh == "c":
            x += (w - larg_caixa) / 2
        elif alinh == "r":
            x += w - larg_caixa
        tb = self.sh.add_textbox(E(x), E(y), E(larg_caixa), E(max(caixa_h, px_f * lh)))
        if nome:
            tb.name = nome
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancora]
        for i, par in enumerate(str(conteudo).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[alinh]
            p.line_spacing = pt(px_f * lh)
            for trecho, tipo in C.segmentos(par):
                if tipo == "dado":
                    trecho = trecho.replace(" ", NBSP)
                self._run(p, trecho, estilo, px_f, cor, tipo, caps)
        return alt, px_f

    def lista(self, x, y, w, itens, px=28, cor_marca=None, lado=12, recuo=36, gap=16, lh=1.4, cor=None):
        cor_marca = cor_marca or C.COR["cinza"]
        yy = y
        for i, it in enumerate(itens):
            if i:
                yy += gap
            self.ret(x, yy + px * lh / 2 - lado / 2 - px * 0.06, lado, lado, cor_marca)
            alt, _ = self.texto(x + recuo, yy, w - recuo, it, "texto", px, cor, lh)
            yy += alt
        return yy - y

    # ---- peças repetidas ------------------------------------------------
    def cabecalho(self, x=MX, y=MT, w=1560, titulo_px=64, sub_campo="subtitulo", titulo_campo="titulo", max_linhas=2):
        s = self.s
        y += self.tracado(x, y) + 24
        if s.get("rotulo"):
            a, _ = self.texto(x, y, w, s["rotulo"], "rotulo", 20, self.c_apoio, 1.3, caps=True)
            y += a + 16
        a, _ = self.texto(x, y, w, s.get(titulo_campo, ""), "display", titulo_px, self.c_txt, 1.03, max_linhas=max_linhas)
        y += a
        if s.get(sub_campo):
            y += 24
            a, _ = self.texto(x, y, min(w, 1200), s[sub_campo], "texto", 28, self.c_apoio, 1.45)
            y += a
        return y

    def fonte(self):
        s = self.s
        if not s.get("fonte"):
            return BASE
        txt = s["fonte"]
        if not C.texto_puro(txt).lower().startswith("fonte"):
            txt = "Fonte: " + txt
        self.texto(MX, FONTE_Y, 1696, txt, "texto", 18, self.c_apoio, 1.4)
        return FONTE_Y - 24

    def rodape(self):
        arq = {"claro": "lg-simbolo.png", "apoio": "lg-simbolo.png", "escuro": "lg-simbolo-negativo.png",
               "azul": "lg-simbolo-mono-branco.png"}[self.fundo]
        self.imagem(C.LOGOS / "png" / arq, MX, 1000, 26.5, 32)
        cor = self.c_apoio
        self.texto(MX + 43, 1000, 1200, self.ctx.get("rodape") or "LG Mobilidade Urbana", "texto", 18, cor, 1.0, h=32, ancora="m")
        self.texto(1608, 1000, 200, f"{self.n}", "rotulo", 18, cor, 1.0, alinh="r", h=32, ancora="m")

    def logo(self, x=MX, y=MT, w=300):
        arq = {"claro": "lg-horizontal.png", "apoio": "lg-horizontal.png", "escuro": "lg-horizontal-negativo.png",
               "azul": "lg-horizontal-mono-branco.png"}[self.fundo]
        pic = self.imagem(C.LOGOS / "png" / arq, x, y, w)
        pic.name = "Logo LG Mobilidade Urbana"


# --------------------------------------------------------------------------
# gráficos nativos
# --------------------------------------------------------------------------
def _sem_fundo(chart):
    cs = chart._chartSpace
    for el in cs.findall(qn("c:spPr")):
        cs.remove(el)
    sp = etree.Element(qn("c:spPr"))
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    cs.find(qn("c:chart")).addnext(sp)


def _fonte_chart(chart, px=21, cor=C.COR["cinza_texto"]):
    chart.font.name = NOME_TEXTO
    chart.font.size = pt(px)
    chart.font.color.rgb = rgb(cor)


def _eixos(chart, s, ini, fim, passo, casas_eixo, horizontal=False):
    va = chart.value_axis
    va.minimum_scale = ini
    va.maximum_scale = fim
    va.major_unit = passo
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(C.COR["borda"])
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.number_format = C.formato_excel(casas_eixo, s.get("prefixo", ""), s.get("sufixo", ""))
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = pt(20)
    va.tick_labels.font.color.rgb = rgb(C.COR["cinza_texto"])
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb(C.COR["cinza"])
    ca.format.line.width = Pt(1)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = pt(21)
    ca.tick_labels.font.color.rgb = rgb(C.COR["grafite"])
    if horizontal:
        # primeira categoria em cima; o eixo de valores cruza no fim e volta para baixo
        ca.reverse_order = True
        cr = va._element.find(qn("c:crosses"))
        if cr is not None:
            cr.set("val", "max")


def grafico_barras(f: Folha, s: dict, x, y, w, h):
    cats = [C.texto_puro(c) for c in s["categorias"]]
    series = s["series"]
    horizontal = s.get("orientacao") == "horizontal"
    fmt = C.formato_excel(s.get("casas"), s.get("prefixo", ""), s.get("sufixo", ""))
    cd = CategoryChartData(number_format=fmt)
    cd.categories = cats
    for j, se in enumerate(series):
        cd.add_series(C.texto_puro(se.get("nome", f"Série {j + 1}")), se["valores"])
    tipo = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = f.sh.add_chart(tipo, E(x), E(y), E(w), E(h), cd)
    gf.name = "Gráfico de barras"
    ch = gf.chart
    ch.has_title = False
    _sem_fundo(ch)
    _fonte_chart(ch)
    todos = [v for se in series for v in se["valores"]]
    ini, fim, passo = C.escala_bonita(min(0, min(todos)), max(0, max(todos)), 4)
    _eixos(ch, s, ini, fim, passo, RH._casas_eixo(passo), horizontal)
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = pt(22)
        ch.legend.font.color.rgb = rgb(C.COR["grafite"])
    plot = ch.plots[0]
    plot.gap_width = 70 if len(series) == 1 else 50
    if len(series) > 1:
        plot.overlap = -8
    if len(cats) * len(series) <= 24:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = pt(22)
        dl.font.bold = True
        dl.font.color.rgb = rgb(C.COR["grafite"])
    destaque = C.indices_destaque(s.get("destaque"))
    for j, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.line.fill.background()
        if len(series) == 1:
            ser.format.fill.fore_color.rgb = rgb(C.COR["azul"] if not destaque else C.COR["cinza_claro"])
            for i in destaque:
                if 0 <= i < len(cats):
                    pnt = ser.points[i]
                    pnt.format.fill.solid()
                    pnt.format.fill.fore_color.rgb = rgb(C.COR["azul"])
        else:
            ser.format.fill.fore_color.rgb = rgb(C.SERIES[j % 4])


def grafico_linha(f: Folha, s: dict, x, y, w, h):
    cats = [C.texto_puro(c) for c in s["categorias"]]
    series = s["series"]
    fmt = C.formato_excel(s.get("casas"), s.get("prefixo", ""), s.get("sufixo", ""))
    cd = CategoryChartData(number_format=fmt)
    cd.categories = cats
    for j, se in enumerate(series):
        cd.add_series(C.texto_puro(se.get("nome", f"Série {j + 1}")), se["valores"])
    gf = f.sh.add_chart(XL_CHART_TYPE.LINE_MARKERS, E(x), E(y), E(w), E(h), cd)
    gf.name = "Gráfico de linha"
    ch = gf.chart
    ch.has_title = False
    _sem_fundo(ch)
    _fonte_chart(ch)
    todos = [v for se in series for v in se["valores"]]
    vmin, vmax = min(todos), max(todos)
    ini, fim, passo = C.escala_bonita(RH._y_min_linha(s, vmin, vmax), vmax, 4)
    _eixos(ch, s, ini, fim, passo, RH._casas_eixo(passo))
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = pt(22)
        ch.legend.font.color.rgb = rgb(C.COR["grafite"])
    n = len(cats)
    destacar = s.get("destacar_ultimo", True)
    for j, ser in enumerate(ch.plots[0].series):
        cor = C.SERIES[j % 4]
        ser.smooth = False
        ser.format.line.color.rgb = rgb(cor)
        ser.format.line.width = E(6 if j == 0 else 4)
        ser.marker.style = XL_MARKER_STYLE.SQUARE
        ser.marker.size = 8 if j == 0 else 6
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb("#FFFFFF")
        ser.marker.format.line.color.rgb = rgb(cor)
        ser.marker.format.line.width = Pt(2)
        vals = series[j]["valores"]
        if j == 0 and destacar:
            mk = ser.points[n - 1].marker
            mk.style = XL_MARKER_STYLE.SQUARE
            mk.size = 12
            mk.format.fill.solid()
            mk.format.fill.fore_color.rgb = rgb(C.COR["verde"])
            mk.format.line.color.rgb = rgb(C.COR["verde"])
        for i, v in enumerate(vals):
            ultimo = i == n - 1
            mostrar = (j == 0 and (n <= 8 or i in (0, n - 1))) or (j > 0 and ultimo)
            if not mostrar:
                continue
            dl = ser.points[i].data_label
            dl.has_text_frame = True
            dl.text_frame.text = C.numero_br(v, s.get("casas"), s.get("prefixo", ""), s.get("sufixo", ""))
            run = dl.text_frame.paragraphs[0].runs[0]
            run.font.size = pt(22)
            run.font.bold = True
            run.font.name = NOME_TEXTO
            run.font.color.rgb = rgb(C.COR["grafite"])
            vale = 0 < i < n - 1 and vals[i - 1] > v and vals[i + 1] > v
            dl.position = XL_LABEL_POSITION.RIGHT if ultimo else (XL_LABEL_POSITION.BELOW if vale else XL_LABEL_POSITION.ABOVE)


def grafico_donut(f: Folha, s: dict, x, y, lado):
    itens = s["itens"]
    cd = CategoryChartData(number_format="0")
    cd.categories = [C.texto_puro(it.get("rotulo")) for it in itens]
    cd.add_series("Participação", [it["valor"] for it in itens])
    gf = f.sh.add_chart(XL_CHART_TYPE.DOUGHNUT, E(x), E(y), E(lado), E(lado), cd)
    gf.name = "Gráfico de rosca"
    ch = gf.chart
    ch.has_title = False
    _sem_fundo(ch)
    _fonte_chart(ch)
    ch.has_legend = False
    plot = ch.plots[0]
    dn = plot._element
    hole = dn.find(qn("c:holeSize"))
    if hole is None:
        hole = etree.SubElement(dn, qn("c:holeSize"))
    hole.set("val", "64")
    cores = RH.cores_fatias(itens)
    ser = plot.series[0]
    ser.format.line.color.rgb = rgb(f.cor_fundo)
    ser.format.line.width = Pt(3)
    for i, cor in enumerate(cores):
        pnt = ser.points[i]
        pnt.format.fill.solid()
        pnt.format.fill.fore_color.rgb = rgb(cor)
        pnt.format.line.color.rgb = rgb(f.cor_fundo)
        pnt.format.line.width = Pt(3)
    # o gráfico ocupa a área com uma pequena margem interna; o texto do centro vai por cima
    centro = s.get("centro") or {}
    cx, cy = x + lado / 2, y + lado / 2
    if centro.get("valor"):
        f.texto(cx - 180, cy - 92 if centro.get("rotulo") else cy - 60, 360, centro["valor"], "display", 104, C.COR["grafite"], 1.0, alinh="c")
    if centro.get("rotulo"):
        f.texto(cx - 150, cy + 24, 300, centro["rotulo"], "texto", 24, C.COR["cinza_texto"], 1.25, alinh="c")


# --------------------------------------------------------------------------
# layouts
# --------------------------------------------------------------------------
def p_capa(f: Folha):
    s = f.s
    f.planta("claro", 1060, 0, 860, 1080, crop=(290 / 1440, 0, 1 - 1150 / 1440, 0))
    f.esmaecer(1058, 0, 200, 1080, C.COR["fundo"])
    ox = 1060
    f.poli([(ox + 300, 1110), (ox + 300, 900), (ox + 460, 740), (ox + 880, 740)], C.COR["grafite"], 12)
    f.poli([(ox + 880, 560), (ox + 660, 560), (ox + 500, 400), (ox + 500, 236)], C.COR["azul"], 12)
    f.quadrado(ox + 300, 900, 28, "#FFFFFF", C.COR["grafite"], 5)
    f.quadrado(ox + 660, 560, 28, "#FFFFFF", C.COR["azul"], 5)
    f.quadrado(ox + 500, 216, 40, C.COR["verde"])
    f.logo()
    w = 960
    meta = " · ".join(v for v in (s.get("data") or f.ctx.get("data"), s.get("local")) if v)
    # altura total para ancorar o bloco embaixo
    titulo_px = 112
    while titulo_px > 112 * 0.8 and len(MED.linhas(puro_nbsp(s["titulo"]), "display", titulo_px, w)) > 4:
        titulo_px *= 0.96
    alt = 40 + 24
    if s.get("rotulo"):
        alt += 26 + 16
    alt += MED.altura(puro_nbsp(s["titulo"]), "display", titulo_px, w, 0.98)
    if s.get("subtitulo"):
        alt += 24 + MED.altura(puro_nbsp(s["subtitulo"]), "texto", 32, 1120, 1.35)
    if meta:
        alt += 48 + 26
    y = 1080 - 96 - alt
    y += f.tracado(MX, y) + 24
    if s.get("rotulo"):
        a, _ = f.texto(MX, y, w, s["rotulo"], "rotulo", 20, f.c_apoio, 1.3, caps=True)
        y += a + 16
    a, _ = f.texto(MX, y, w, s["titulo"], "display", titulo_px, f.c_txt, 0.98)
    y += a
    if s.get("subtitulo"):
        y += 24
        a, _ = f.texto(MX, y, 1120, s["subtitulo"], "texto", 32, f.c_apoio, 1.35)
        y += a
    if meta:
        f.texto(MX, y + 48, w, meta, "rotulo", 20, f.c_apoio, 1.3, caps=True)


def p_secao(f: Folha):
    s = f.s
    f.planta("escuro", 0, 0, 1920, 1080, crop=(0, 0.125, 0, 0.125))
    f.poli([(-20, 300), (1180, 300), (1340, 140), (1940, 140)], C.COR["azul_claro"], 12)
    f.quadrado(1180, 300, 28, C.COR["grafite"], C.COR["azul_claro"], 5)
    w = 1560
    num = s.get("numero")
    alt = 0
    if num not in (None, ""):
        alt += 172 + 40
    alt += 64
    alt += MED.altura(puro_nbsp(s["titulo"]), "display", 96, w, 1.0)
    if s.get("descricao"):
        alt += 24 + MED.altura(puro_nbsp(s["descricao"]), "texto", 32, 1120, 1.35)
    y = BASE - alt
    if num not in (None, ""):
        f.texto(MX, y, 800, str(num), "display", 200, C.COR["azul_claro"], 0.86)
        y += 172 + 40
    y += f.tracado(MX, y) + 24
    a, _ = f.texto(MX, y, w, s["titulo"], "display", 96, f.c_txt, 1.0)
    y += a
    if s.get("descricao"):
        f.texto(MX, y + 24, 1120, s["descricao"], "texto", 32, f.c_apoio, 1.35)


def p_afirmacao(f: Folha):
    s = f.s
    w = 1560
    tpx = 96
    while tpx > 96 * 0.8 and len(MED.linhas(puro_nbsp(s["titulo"]), "display", tpx, w)) > 4:
        tpx *= 0.96
    alt = 64 + (42 if s.get("rotulo") else 0) + MED.altura(puro_nbsp(s["titulo"]), "display", tpx, w, 1.0)
    if s.get("texto"):
        alt += 24 + MED.altura(puro_nbsp(s["texto"]), "texto", 32, 1120, 1.35)
    limite = f.fonte()
    y = MT + (limite - MT - alt) / 2
    y += f.tracado(MX, y) + 24
    if s.get("rotulo"):
        a, _ = f.texto(MX, y, w, s["rotulo"], "rotulo", 20, f.c_apoio, 1.3, caps=True)
        y += a + 16
    a, _ = f.texto(MX, y, w, s["titulo"], "display", tpx, f.c_txt, 1.0)
    y += a
    if s.get("texto"):
        f.texto(MX, y + 24, 1120, s["texto"], "texto", 32, f.c_apoio, 1.35)


def p_kpis(f: Folha):
    s = f.s
    itens = s["itens"]
    n = len(itens)
    topo = f.cabecalho() + 48
    base = f.fonte()
    colw = (1696 - 48 * (n - 1)) / n
    kv = {2: 200, 3: 168, 4: 136}.get(n, 136)
    larg_txt = colw - 32
    def so_dado(v):
        return bool(C.RE_DADO.fullmatch(str(v or "").strip()))
    for it in itens:
        if so_dado(it.get("valor")):
            continue
        txt = C.texto_puro(it.get("valor"))
        wv = MED.largura(txt, "display", kv)
        if wv > larg_txt:
            kv *= larg_txt / wv * 0.97
    alturas = []
    for it in itens:
        a = kv * 0.92 + 8
        if it.get("rotulo"):
            a += 24 + MED.altura(puro_nbsp(it["rotulo"]), "texto-sb", 28, larg_txt, 1.3)
        if it.get("legenda"):
            a += 8 + MED.altura(puro_nbsp(it["legenda"]), "texto", 24, larg_txt, 1.45)
        alturas.append(a)
    bloco = max(alturas)
    y0 = topo + max(0, (base - topo - bloco) / 2)
    for i, it in enumerate(itens):
        x = MX + i * (colw + 48)
        f.ret(x, y0, 1, bloco, C.COR["cinza"])
        y = y0 + 8
        cor = C.COR["azul"] if it.get("destaque") else f.c_txt
        if so_dado(it.get("valor")):
            f.texto(x + 32, y + kv * 0.92 - 52, larg_txt, it.get("valor"), "texto", 36, cor, 1.2)
            y += kv * 0.92
        else:
            a, _ = f.texto(x + 32, y, larg_txt, it.get("valor"), "display", kv, cor, 0.92)
            y += a
        if it.get("rotulo"):
            a, _ = f.texto(x + 32, y + 24, larg_txt, it["rotulo"], "texto-sb", 28, f.c_txt, 1.3)
            y += 24 + a
        if it.get("legenda"):
            f.texto(x + 32, y + 8, larg_txt, it["legenda"], "texto", 24, f.c_apoio, 1.45)


def p_grafico(f: Folha, tipo: str):
    topo = f.cabecalho() + 32
    base = f.fonte()
    if tipo == "barras":
        grafico_barras(f, f.s, MX - 8, topo, 1712, base - topo)
    else:
        grafico_linha(f, f.s, MX - 8, topo, 1712, base - topo)


def p_donut(f: Folha):
    s = f.s
    topo = f.cabecalho() + 48
    base = f.fonte()
    lado = min(600, base - topo)
    y0 = topo + (base - topo - lado) / 2
    grafico_donut(f, s, MX, y0, lado)
    cores = RH.cores_fatias(s["itens"])
    vals = RH.valores_donut(s)
    xl = MX + lado + 112
    wl = 1808 - xl
    linha_h = 72
    alt = linha_h * len(s["itens"])
    y = topo + (base - topo - alt) / 2
    for i, it in enumerate(s["itens"]):
        f.ret(xl, y, wl, 1, C.COR["borda"])
        f.ret(xl, y + linha_h / 2 - 12, 24, 24, cores[i])
        f.texto(xl + 48, y, wl - 48 - 160, it.get("rotulo"), "texto", 28, f.c_txt, 1.3, h=linha_h, ancora="m")
        f.texto(1808 - 160, y, 160, vals[i], "texto-sb", 28, f.c_txt, 1.3, alinh="r", h=linha_h, ancora="m")
        y += linha_h
    f.ret(xl, y, wl, 1, C.COR["borda"])


def _celula(cel, txt, estilo, px, cor, alinh, fundo_hex, negrito=False, caps=False, spc=0.0):
    cel.fill.solid()
    cel.fill.fore_color.rgb = rgb(fundo_hex)
    cel.margin_left = cel.margin_right = E(24)
    cel.margin_top = cel.margin_bottom = E(12)
    cel.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cel.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT if alinh == "r" else PP_ALIGN.LEFT
    for trecho, _tipo in C.segmentos(txt):
        r = p.add_run()
        r.text = trecho
        nome = FONTES["texto-sb" if negrito else estilo][0]
        r.font.name = nome
        r.font.size = pt(px)
        r.font.color.rgb = rgb(cor)
        rpr = r._r.get_or_add_rPr()
        if caps:
            rpr.set("cap", "all")
        if spc:
            rpr.set("spc", str(int(round(spc * px * 0.5 * 100))))


def _bordas(cel, baixo=None, esquerda=None, larg_esq=0):
    tcPr = cel._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)

    def ln(tag, cor, larg):
        el = etree.Element(qn(tag), w=str(int(larg)), cap="flat", cmpd="sng", algn="ctr")
        if cor:
            sf = etree.SubElement(el, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=cor.lstrip("#").upper())
        else:
            etree.SubElement(el, qn("a:noFill"))
        return el
    novos = [ln("a:lnL", esquerda, larg_esq * EMU_PX if esquerda else 0), ln("a:lnR", None, 0),
             ln("a:lnT", None, 0), ln("a:lnB", baixo, EMU_PX if baixo else 0)]
    for el in reversed(novos):
        tcPr.insert(0, el)


def p_tabela(f: Folha):
    s = f.s
    topo = f.cabecalho() + 48
    base = f.fonte()
    cols, linhas = s["colunas"], s["linhas"]
    nums = RH.colunas_numericas(s)
    dest = s.get("destaque_linha")
    # largura proporcional ao maior texto de cada coluna
    larg = []
    for j, c in enumerate(cols):
        w = MED.largura(C.texto_puro(c), "rotulo", 18, caps=True)
        for ln_ in linhas:
            w = max(w, MED.largura(C.texto_puro(RH.texto_celula(ln_[j])), "texto-sb", 24))
        larg.append(w + 48)
    total = sum(larg)
    larg = [w / total * 1696 for w in larg]
    hh, hl_ = 56, 64
    gt = f.sh.add_table(len(linhas) + 1, len(cols), E(MX), E(topo), E(1696), E(hh + hl_ * len(linhas)))
    gt.name = "Tabela"
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tbl.first_row = True
    tbl.horz_banding = False
    for j, w in enumerate(larg):
        tbl.columns[j].width = E(w)
    tbl.rows[0].height = E(hh)
    for j, c in enumerate(cols):
        cel = tbl.cell(0, j)
        _celula(cel, C.texto_puro(c), "rotulo", 18, C.COR["cinza_texto"], "r" if nums[j] else "l", C.COR["apoio"], caps=True, spc=0.1)
        _bordas(cel)
    for i, ln_ in enumerate(linhas, 1):
        tbl.rows[i].height = E(hl_)
        for j, c in enumerate(ln_):
            cel = tbl.cell(i, j)
            negrito = dest == i - 1
            _celula(cel, RH.texto_celula(c), "texto", 24, C.COR["grafite"], "r" if nums[j] else "l", "#FFFFFF", negrito)
            _bordas(cel, baixo=C.COR["borda"] if i < len(linhas) else None,
                    esquerda=C.COR["azul"] if (negrito and j == 0) else None, larg_esq=6)
    if topo + hh + hl_ * len(linhas) > base:
        f.avisos.append(f"slide {f.n:02d}: tabela passa do limite inferior; reduza linhas")


def _altura_lado(b: dict, w: float) -> float:
    a = 0
    if b.get("rotulo"):
        a += 26 + 16
    if b.get("titulo"):
        a += MED.altura(puro_nbsp(b["titulo"]), "h3", 34, w, 1.12)
    a += altura_lista(w, [C.texto_puro(i) for i in b.get("itens", [])], 28) + 16
    return a


def p_comparativo(f: Folha):
    s = f.s
    topo = f.cabecalho() + 48
    base = f.fonte()
    lw = (1696 - 160) / 2
    wi = lw - 96
    alt = max(_altura_lado(s["esquerda"], wi), _altura_lado(s["direita"], wi)) + 96
    y0 = topo + max(0, (base - topo - alt) / 2)
    for k, (b, x) in enumerate(((s["esquerda"], MX), (s["direita"], MX + lw + 160))):
        if k == 0:
            f.ret(x, y0, lw, alt, C.COR["apoio"])
        else:
            f.ret(x, y0, lw, alt, "#FFFFFF")
            f.ret(x, y0, lw, 6, C.COR["azul"])
        y = y0 + 48
        if b.get("rotulo"):
            a, _ = f.texto(x + 48, y, wi, b["rotulo"], "rotulo", 20, C.COR["azul"] if k else C.COR["cinza_texto"], 1.3, caps=True)
            y += a + 16
        if b.get("titulo"):
            a, _ = f.texto(x + 48, y, wi, b["titulo"], "h3", 34, f.c_txt, 1.12)
            y += a
        f.lista(x + 48, y + 16, wi, b.get("itens", []), 28, C.COR["azul"] if k else C.COR["cinza"])
    cy = y0 + alt / 2
    xa = MX + lw
    f.poli([(xa, cy + 24), (xa + 52, cy + 24), (xa + 88, cy - 12), (xa + 160, cy - 12)], C.COR["azul"], 8)


def p_tempo(f: Folha):
    s = f.s
    marcos = s["marcos"]
    n = len(marcos)
    topo = f.cabecalho() + 48
    base = f.fonte()
    colw = (1696 - 32 * (n - 1)) / n
    tx_alt = 0
    for m in marcos:
        a = 24 + MED.altura(puro_nbsp(m.get("titulo")), "h3", 28, colw - 8, 1.12)
        if m.get("texto"):
            a += 8 + MED.altura(puro_nbsp(m["texto"]), "texto", 24, colw - 8, 1.45)
        tx_alt = max(tx_alt, a)
    bloco = 24 + 120 + 56 + tx_alt
    y0 = topo + max(0, (base - topo - bloco) / 2)
    ly = y0 + 24 + 120 + 28
    f.ret(0, ly - 3, 1920, 6, C.COR["grafite"])
    atual = s.get("atual")
    chegada = s.get("chegada", True)
    for i, m in enumerate(marcos):
        x = MX + i * (colw + 32)
        f.texto(x, y0 + 24, colw, m.get("data"), "display", 44, C.COR["azul"], 1.0, h=104, ancora="b")
        if chegada and i == n - 1:
            f.quadrado(x + 18, ly, 36, C.COR["verde"])
        elif atual == i:
            f.quadrado(x + 14, ly, 28, C.COR["azul"])
        else:
            f.quadrado(x + 14, ly, 28, "#FFFFFF", C.COR["grafite"], 4)
        y = y0 + 24 + 120 + 56 + 24
        a, _ = f.texto(x, y, colw - 8, m.get("titulo"), "h3", 28, f.c_txt, 1.12)
        if m.get("texto"):
            f.texto(x, y + a + 8, colw - 8, m["texto"], "texto", 24, f.c_apoio, 1.45)


def p_rede(f: Folha):
    s = f.s
    ox = 760
    f.planta("claro", ox, 0, 1160, 1080, crop=(280 / 1440, 0, 0, 0))
    p, u, x0, y0 = RH.geometria_rede(s)

    def P(gx, gy):
        x, y = p(gx, gy)
        return ox + x, y
    nos = {nd["id"]: nd for nd in s["nos"]}
    for j, ln_ in enumerate(s["linhas"]):
        cor = C.COR_LINHA_REDE.get(ln_.get("cor", ""), C.SERIES[j % 3])
        pts = C.trajeto([P(nos[i]["x"], nos[i]["y"]) for i in ln_["nos"]])
        shp = f.poli(pts, cor, 12)
        shp.name = f"Linha {C.texto_puro(ln_.get('nome'))}"
    for nd in s["nos"]:
        x, y = P(nd["x"], nd["y"])
        tipo = nd.get("tipo", "parada")
        if tipo == "terminal":
            f.quadrado(x, y, 40, "#FFFFFF", C.COR["grafite"], 6)
            r = 20
        elif tipo == "chegada":
            f.quadrado(x, y, 36, C.COR["verde"])
            r = 18
        else:
            f.quadrado(x, y, 24, "#FFFFFF", C.COR["grafite"], 5)
            r = 12
        if nd.get("rotulo"):
            estilo = "texto-sb" if tipo in ("terminal", "chegada") else "texto"
            txt = C.texto_puro(nd["rotulo"])
            w = MED.largura(txt, estilo, 22) + 12
            pos = nd.get("rotulo_pos", "direita")
            if pos == "esquerda":
                bx, by, al = x - r - 14 - w, y - 16, "r"
            elif pos == "acima":
                bx, by, al = x - w / 2, y - r - 16 - 26, "c"
            elif pos == "abaixo":
                bx, by, al = x - w / 2, y + r + 8, "c"
            else:
                bx, by, al = x + r + 14, y - 16, "l"
            f.texto(bx, by, w, txt, estilo, 22, C.COR["grafite"], 1.3, alinh=al)
    if s.get("escala"):
        ex, ey = ox + x0, y0 + 8 * u + 88
        f.poli([(ex, ey - 10), (ex, ey), (ex + 2 * u, ey), (ex + 2 * u, ey - 10)], C.COR["grafite"], 3)
        f.texto(ex + 2 * u + 16, ey - 20, 400, C.texto_puro(s["escala"]), "texto", 20, C.COR["grafite"], 1.2)
    y = f.cabecalho(w=600, titulo_px=64, sub_campo="texto", max_linhas=5)
    if s["linhas"]:
        y += 40
        for j, ln_ in enumerate(s["linhas"]):
            cor = C.COR_LINHA_REDE.get(ln_.get("cor", ""), C.SERIES[j % 3])
            f.ret(MX, y + 16 - 6, 48, 12, cor)
            a, _ = f.texto(MX + 64, y, 536, ln_.get("nome"), "texto", 24, f.c_txt, 1.35)
            y += max(a, 32) + 16
    f.fonte()


def p_citacao(f: Folha):
    s = f.s
    txt = s["texto"].strip()
    if not txt.startswith(("“", '"')):
        txt = f"“{txt}”"
    w = 1500
    extra = " · ".join(v for v in (s.get("cargo"), s.get("fonte")) if v)
    alt = 64 + MED.altura(puro_nbsp(txt), "display", 64, w, 1.08) + 48 + 29
    if extra:
        alt += 8 + 35
    y = MT + (BASE - MT - alt) / 2
    y += f.tracado(MX, y) + 24
    a, _ = f.texto(MX, y, w, txt, "display", 64, f.c_txt, 1.08)
    y += a + 48
    f.texto(MX, y, w, s["autor"], "rotulo", 22, f.c_txt, 1.3, caps=True)
    if extra:
        f.texto(MX, y + 29 + 8, w, extra, "texto", 24, f.c_apoio, 1.45)


def p_pilares(f: Folha):
    s = f.s
    lw = (1696 - 112) * 5 / 12
    rx = MX + lw + 112
    rw = 1808 - rx
    base = f.fonte()
    # coluna da esquerda centrada na altura
    alt_e = 64 + (42 if s.get("rotulo") else 0) + MED.altura(puro_nbsp(s["titulo"]), "display", 64, lw, 1.03)
    if s.get("texto"):
        alt_e += 24 + MED.altura(puro_nbsp(s["texto"]), "texto", 28, lw, 1.45)
    f.cabecalho(x=MX, y=MT + (base - MT - alt_e) / 2, w=lw, sub_campo="texto", max_linhas=4)
    itens = s["itens"]
    alturas = []
    for it in itens:
        a = 32 + MED.altura(puro_nbsp(it.get("titulo")), "h3", 34, rw - 72, 1.12)
        if it.get("texto"):
            a += 8 + MED.altura(puro_nbsp(it["texto"]), "texto", 24, rw - 72, 1.45)
        alturas.append(a + 32)
    total = sum(alturas)
    y = MT + (base - MT - total) / 2
    for i, it in enumerate(itens):
        f.ret(rx, y, rw, 1, C.COR["borda"])
        f.ret(rx, y + 32 + 17 - 10, 20, 20, C.COR["azul"])
        a, _ = f.texto(rx + 72, y + 32, rw - 72, it.get("titulo"), "h3", 34, f.c_txt, 1.12)
        if it.get("texto"):
            f.texto(rx + 72, y + 32 + a + 8, rw - 72, it["texto"], "texto", 24, f.c_apoio, 1.45)
        y += alturas[i]
    f.ret(rx, y, rw, 1, C.COR["borda"])


def p_processo(f: Folha):
    s = f.s
    etapas = s["etapas"]
    k = len(etapas)
    topo = f.cabecalho() + 48
    base = f.fonte()
    cw, niveis, pontos, marcos = RH.geometria_processo(k)
    yb = topo + max(0, (base - topo - 480) / 2)
    f.poli([(MX + x, yb + y) for x, y in pontos], C.COR["grafite"], 6)
    for i, (mx, my) in enumerate(marcos):
        if i == k - 1:
            f.quadrado(MX + mx + 18, yb + my, 36, C.COR["verde"])
        else:
            f.quadrado(MX + mx + 14, yb + my, 28, "#FFFFFF", C.COR["grafite"], 5)
        f.texto(MX + mx, yb + my - 96, cw - 48, f"{i + 1}", "display", 56, C.COR["azul"], 1.0)
        et = etapas[i]
        a, _ = f.texto(MX + mx, yb + my + 40, cw - 56, et.get("titulo"), "h3", 30, f.c_txt, 1.12)
        if et.get("texto"):
            f.texto(MX + mx, yb + my + 40 + a + 8, cw - 56, et["texto"], "texto", 24, f.c_apoio, 1.45)


def p_topicos(f: Folha):
    s = f.s
    topo = f.cabecalho() + 48
    f.fonte()
    nota = s.get("nota")
    lw = (1696 - 96) * 8 / 12 if nota else 1696
    y = topo
    for it in s["itens"]:
        f.ret(MX, y + 28 * 1.4 / 2 - 7 - 1, 14, 14, C.COR["azul"])
        a, _ = f.texto(MX + 44, y, lw - 44, it, "texto", 28, f.c_txt, 1.4)
        y += a + 24
    if nota:
        nx = MX + lw + 96
        nw = 1808 - nx
        alt = 80 + MED.altura(puro_nbsp(nota.get("texto")), "texto", 28, nw - 80, 1.45)
        if nota.get("rotulo"):
            alt += 26 + 16
        f.ret(nx, topo, nw, alt, "#FFFFFF")
        f.ret(nx, topo, nw, 6, C.COR["azul"])
        yy = topo + 40
        if nota.get("rotulo"):
            a, _ = f.texto(nx + 40, yy, nw - 80, nota["rotulo"], "rotulo", 20, C.COR["cinza_texto"], 1.3, caps=True)
            yy += a + 16
        f.texto(nx + 40, yy, nw - 80, nota.get("texto"), "texto", 28, f.c_txt, 1.45)


def p_encerramento(f: Folha):
    s = f.s
    traco = "#FFFFFF" if f.fundo == "azul" else C.COR["azul_claro"]
    for pts in ([(-20, 150), (900, 150), (1000, 50), (1940, 50)],
                [(1720, 1100), (1720, 700), (1820, 600), (1940, 600)],
                [(1940, 340), (1560, 340), (1460, 440), (1460, 1100)]):
        f.poli(pts, traco, 12, alpha=0.14)
    for cx, cy in ((900, 150), (1720, 700), (1560, 340)):
        f.quadrado(cx, cy, 20, None, traco, 3, alpha=0.5)
    f.logo()
    w = 1220
    contato = s.get("contato") or []
    alt = MED.altura(puro_nbsp(s["titulo"]), "display", 96, w, 1.0)
    if s.get("rotulo"):
        alt += 42
    if s.get("texto"):
        alt += 24 + MED.altura(puro_nbsp(s["texto"]), "texto", 32, 1120, 1.35)
    if s.get("proximo_passo"):
        alt += 48 + 28 * 1.2 + 40
    if contato:
        alt += 48 + 35
    y = 1080 - 96 - alt
    if s.get("rotulo"):
        a, _ = f.texto(MX, y, w, s["rotulo"], "rotulo", 20, f.c_apoio, 1.3, caps=True)
        y += a + 16
    a, _ = f.texto(MX, y, w, s["titulo"], "display", 96, f.c_txt, 1.0)
    y += a
    if s.get("texto"):
        a, _ = f.texto(MX, y + 24, 1120, s["texto"], "texto", 32, f.c_apoio, 1.35)
        y += 24 + a
    if s.get("proximo_passo"):
        txt = C.texto_puro(s["proximo_passo"])
        bw = MED.largura(txt, "h3", 28) + 64 + 8
        bh = 28 * 1.2 + 40
        shp = f.ret(MX, y + 48, bw, bh, C.COR["verde"], forma=MSO_SHAPE.ROUNDED_RECTANGLE)
        shp.adjustments[0] = 4 / bh
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = E(32)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        f._run(p, txt, "h3", 28, C.COR["grafite"])
        y += 48 + bh
    if contato:
        x = MX
        for c in contato:
            wc = MED.largura(C.texto_puro(c), "texto-sb" if C.RE_DADO.search(c) else "texto", 24) + 12
            f.texto(x, y + 48, wc, c, "texto", 24, "#EAEDF0" if f.fundo == "azul" else C.COR["cinza_claro"], 1.45)
            x += wc + 48


PINTAR = {
    "capa": p_capa, "secao": p_secao, "afirmacao": p_afirmacao, "kpis": p_kpis,
    "grafico-barras": lambda f: p_grafico(f, "barras"), "grafico-linha": lambda f: p_grafico(f, "linha"),
    "grafico-donut": p_donut, "tabela": p_tabela, "comparativo": p_comparativo, "linha-do-tempo": p_tempo,
    "rede": p_rede, "citacao": p_citacao, "pilares": p_pilares, "processo": p_processo,
    "topicos": p_topicos, "encerramento": p_encerramento,
}


def _tema(prs):
    """Troca as fontes e as cores do tema para as da LG (texto novo já nasce na marca)."""
    try:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        parte = prs.slide_master.part.part_related_by(RT.THEME)
        xml = parte.blob.decode("utf-8")
        xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*"', r'\1Archivo SemiCondensed ExtraBold"', xml)
        xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*"', r'\1IBM Plex Sans"', xml)
        cores = {"dk1": "22262B", "lt1": "FFFFFF", "dk2": "2C3138", "lt2": "F9FAFB", "accent1": "1266D3",
                 "accent2": "22262B", "accent3": "6F7D90", "accent4": "AEB8C5", "accent5": "21C45D",
                 "accent6": "2C7FED", "hlink": "1266D3", "folHlink": "105BBC"}
        for nome, val in cores.items():
            xml = re.sub(rf"(<a:{nome}>)\s*<a:(?:srgbClr|sysClr)[^>]*/>", rf'\1<a:srgbClr val="{val}"/>', xml)
        parte._blob = xml.encode("utf-8")
    except Exception as e:  # tema é conforto, não requisito
        print("AVISO: não consegui ajustar o tema do PPTX:", e)


def fontes_instaladas() -> list[str]:
    pastas = []
    if os.name == "nt":
        pastas += [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts",
                   Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft/Windows/Fonts"]
    else:
        pastas += [Path.home() / "Library/Fonts", Path("/Library/Fonts"), Path.home() / ".fonts",
                   Path.home() / ".local/share/fonts", Path("/usr/share/fonts")]
    faltam = []
    for arq in sorted({v[2] for v in FONTES.values()}):
        if not any((p / arq).exists() for p in pastas if p and p.exists()):
            faltam.append(arq)
    return faltam


def render(deck: dict, saida: Path) -> tuple[Path, list[str]]:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    _tema(prs)
    meta = deck.get("meta", {})
    prs.core_properties.title = C.texto_puro(meta.get("titulo") or "Apresentação LG Mobilidade Urbana")
    prs.core_properties.author = "LG Mobilidade Urbana"
    prs.core_properties.language = "pt-BR"
    avisos = []
    for n, s in enumerate(deck["slides"], 1):
        f = Folha(prs, s, n, meta)
        PINTAR[s["layout"]](f)
        if s["layout"] not in RH.SEM_RODAPE:
            f.rodape()
        if s.get("notas"):
            f.sl.notes_slide.notes_text_frame.text = C.texto_puro(s["notas"])
        avisos += f.avisos
    saida.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(saida))
    return saida, avisos


def main(argv=None) -> int:
    C.utf8_console()
    ap = argparse.ArgumentParser(description="Gera o PPTX editável da apresentação LG a partir do deck.json.")
    ap.add_argument("deck", type=Path)
    ap.add_argument("-o", "--saida", type=Path)
    a = ap.parse_args(argv)
    deck = C.ler_deck(a.deck)
    erros, avs = C.validar(deck)
    for av in avs:
        print("AVISO:", av)
    if erros:
        for e in erros:
            print("ERRO:", e)
        return 2
    saida, avisos = render(deck, a.saida or a.deck.with_suffix(".pptx"))
    for av in avisos:
        print("AVISO PPTX:", av)
    faltam = fontes_instaladas()
    if faltam:
        print("AVISO: fontes da marca não instaladas neste computador:", ", ".join(faltam),
              "(instale os TTF de assets/fonts/ttf para ver o PPTX como foi desenhado)")
    print(f"PPTX: {saida.resolve()} ({saida.stat().st_size // 1024} KB, {len(deck['slides'])} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
